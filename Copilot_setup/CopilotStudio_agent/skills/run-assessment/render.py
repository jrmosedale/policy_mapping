#!/usr/bin/env python3
"""
render.py - build a Met Office research -> policy relevance assessment as PDF and/or Word
from a single YAML config, so every field uses the same method, scoring and layout.

Usage:
    python render.py assessment_config.yaml --format both --outdir outputs
    python render.py assessment_config.yaml --format pdf
    python render.py assessment_config.yaml --format docx

Authoring markup usable in any text field:  **bold**   *italic*   [[label|https://url]]
Edit headings, section order, colours and labels in the YAML, not here.
"""
import argparse, os, re, sys, yaml

# ----------------------------------------------------------------------------- helpers
def parse_markup(text, bold=False, italic=False, url=None):
    """Recursively turn **bold**/*italic*/[[label|url]] into a flat list of run specs."""
    out, pos = [], 0
    pat = r'(\*\*[\s\S]+?\*\*|\*[^*]+?\*|\[\[[^\]]+\]\])'
    for m in re.finditer(pat, text):
        if m.start() > pos:
            out.append(dict(text=text[pos:m.start()], bold=bold, italic=italic, url=url))
        tok = m.group(0)
        if tok.startswith('**'):
            out += parse_markup(tok[2:-2], True, italic, url)
        elif tok.startswith('[['):
            i = tok.index('|'); out += parse_markup(tok[2:i], bold, italic, tok[i+1:-2])
        else:
            out += parse_markup(tok[1:-1], bold, True, url)
        pos = m.end()
    if pos < len(text):
        out.append(dict(text=text[pos:], bold=bold, italic=italic, url=url))
    return out

def band_of(relevance):
    return {3: 'H', 2: 'M', 1: 'L'}[int(relevance)]

DEFAULT_THEME = dict(
    teal='0B5563', teal_dk='08414C', ink='1F2A33', muted='5A6B73', rule='C5D0D4',
    high_bg='CDE9D2', high_chip='2E7D43', med_bg='FBE3B8', med_chip='C9871B',
    low_bg='DCE4E9', low_chip='5A6B73', status_current='0B5563', status_opportunity='B07012',
    link='10497A', kp_fill='EAF2F3', alt_tint='F2F6F7', font='Arial')

def load_config(path):
    with open(path, 'r', encoding='utf-8') as fh:
        cfg = yaml.safe_load(fh)
    th = dict(DEFAULT_THEME); th.update(cfg.get('theme', {})); cfg['theme'] = th
    cfg.setdefault('section_order',
                   ['key_points', 'summary', 'policy_table', 'indicators', 'methods', 'references'])
    return cfg

# =============================================================================  PDF
def build_pdf(cfg, out_path):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT, TA_CENTER
    from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph,
                                    Spacer, Table, TableStyle, HRFlowable, KeepTogether)

    T = cfg['theme']
    C = {k: colors.HexColor('#' + v) for k, v in T.items() if re.fullmatch(r'[0-9A-Fa-f]{6}', str(v))}
    LINK = '#' + T['link']
    BAND_BG = {'H': C['high_bg'], 'M': C['med_bg'], 'L': C['low_bg']}
    BAND_CHIP = {'H': C['high_chip'], 'M': C['med_chip'], 'L': C['low_chip']}
    _sa = cfg.get('status_axis') or {}
    status_label = _sa.get('label', 'Status')
    _sv = _sa.get('values') or {'Current': T['status_current'], 'Opportunity': T['status_opportunity']}
    STATUS_BG = {k: colors.HexColor('#' + v) for k, v in _sv.items()}
    labels = cfg.get('labels', {})
    rel_lbl = {int(k): v for k, v in cfg.get('scoring', {}).get('relevance',
                {3: 'High', 2: 'Medium', 1: 'Low'}).items()}

    def esc(s):
        return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    def mu(text):
        """mini-markup -> reportlab inline markup."""
        s = esc(text)
        s = re.sub(r'\[\[([^\]|]+)\|([^\]]+)\]\]',
                   lambda m: "<a href='%s' color='%s'><u>%s</u></a>" % (m.group(2), LINK, m.group(1)), s)
        s = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', s)
        s = re.sub(r'\*([^*]+?)\*', r'<i>\1</i>', s)
        return s

    ss = getSampleStyleSheet()
    def st(name, **kw):
        return ParagraphStyle(name, parent=kw.pop('parent', ss['Normal']), **kw)
    INK = C['ink']; TEAL = C['teal']; TEAL_DK = C['teal_dk']; MUTED = C['muted']; RULE = C['rule']
    H1 = st('H1', fontName='Helvetica-Bold', fontSize=13, leading=16, textColor=TEAL_DK,
            spaceBefore=10, spaceAfter=5)
    BODY = st('Body', fontName='Helvetica', fontSize=9.3, leading=13.2, textColor=INK, spaceAfter=5)
    BODY_SM = st('BodySm', fontName='Helvetica', fontSize=8.5, leading=12, textColor=INK, spaceAfter=4)
    BULLET = st('Bullet', parent=BODY, leftIndent=11, spaceAfter=3)
    REFP = st('Ref', fontName='Helvetica', fontSize=8.3, leading=11.5, textColor=INK, spaceAfter=4)
    TH = st('TH', fontName='Helvetica-Bold', fontSize=8.6, leading=10.5, textColor=colors.white)
    TH_C = st('THc', parent=TH, alignment=TA_CENTER)
    C_REF = st('Cref', fontName='Helvetica-Bold', fontSize=8.4, leading=10.5, textColor=INK)
    C_NAME = st('Cname', fontName='Helvetica-Bold', fontSize=8.4, leading=10.6, textColor=INK)
    C_BODY = st('Cbody', fontName='Helvetica', fontSize=8.3, leading=11.0, textColor=INK)
    CHIP = st('Chip', fontName='Helvetica-Bold', fontSize=8.6, leading=10, textColor=colors.white,
              alignment=TA_CENTER)
    KP_TITLE = st('kpTitle', fontName='Helvetica-Bold', fontSize=10.5, leading=13, textColor=TEAL_DK,
                  spaceAfter=4)
    KP_ITEM = st('kpItem', fontName='Helvetica', fontSize=8.8, leading=12.2, textColor=INK,
                 leftIndent=10, spaceAfter=3)
    IND_TH = st('iTH', fontName='Helvetica-Bold', fontSize=8.6, leading=10.5, textColor=colors.white)
    IND_REF = st('iRef', fontName='Helvetica-Bold', fontSize=8.3, leading=10.4, textColor=INK)
    IND_BODY = st('iBody', fontName='Helvetica', fontSize=8.3, leading=10.8, textColor=INK)
    SUBH = st('subH', parent=BODY, fontName='Helvetica-Bold', textColor=TEAL_DK, spaceAfter=3)

    PAGE = landscape(A4); LM = RM = 14 * mm; TMARGIN = 30 * mm; BM = 13 * mm
    meta = cfg.get('meta', {})
    title = meta.get('title', 'Assessment')
    subtitle = meta.get('subtitle', '')
    prepared = meta.get('prepared', '')
    footer_txt = meta.get('footer', '')

    def header_footer(canvas, doc):
        canvas.saveState(); w, h = PAGE
        canvas.setFillColor(TEAL); canvas.rect(0, h - 24 * mm, w, 24 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.white); canvas.setFont('Helvetica-Bold', 14)
        canvas.drawString(LM, h - 12 * mm, title)
        canvas.setFont('Helvetica', 9); canvas.setFillColor(colors.HexColor('#D6E6E9'))
        canvas.drawString(LM, h - 17.5 * mm, subtitle)
        canvas.setFont('Helvetica', 8); canvas.drawRightString(w - RM, h - 17.5 * mm, prepared)
        canvas.setStrokeColor(RULE); canvas.setLineWidth(0.5)
        canvas.line(LM, BM + 4 * mm, w - RM, BM + 4 * mm)
        canvas.setFillColor(MUTED); canvas.setFont('Helvetica', 7.5)
        canvas.drawString(LM, BM, footer_txt)
        canvas.drawRightString(w - RM, BM, 'Page %d' % doc.page)
        canvas.restoreState()

    doc = BaseDocTemplate(out_path, pagesize=PAGE, leftMargin=LM, rightMargin=RM,
                          topMargin=TMARGIN, bottomMargin=BM + 8 * mm, title=title)
    frame = Frame(LM, BM + 8 * mm, PAGE[0] - LM - RM, PAGE[1] - TMARGIN - (BM + 8 * mm),
                  id='main', leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates([PageTemplate(id='all', frames=[frame], onPage=header_footer)])
    USABLE = PAGE[0] - LM - RM

    # ---- section builders -------------------------------------------------
    def sec_key_points():
        items = cfg.get('key_points', [])
        if not items:
            return []
        flow = [Paragraph(labels.get('key_points', 'Key points'), KP_TITLE)]
        for it in items:
            flow.append(Paragraph('\u2022&nbsp;&nbsp;' + mu(it), KP_ITEM))
        box = Table([[flow]], colWidths=[USABLE])
        box.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), C['kp_fill']),
            ('BOX', (0, 0), (-1, -1), 0.8, TEAL),
            ('LINEBEFORE', (0, 0), (0, -1), 3.2, TEAL),
            ('LEFTPADDING', (0, 0), (-1, -1), 12), ('RIGHTPADDING', (0, 0), (-1, -1), 12),
            ('TOPPADDING', (0, 0), (-1, -1), 9), ('BOTTOMPADDING', (0, 0), (-1, -1), 7)]))
        return [box, Spacer(1, 9)]

    def sec_summary():
        s = cfg.get('summary', {})
        out = [Paragraph(labels.get('summary', 'Summary'), H1)]
        if s.get('intro'):
            out.append(Paragraph(mu(s['intro']), BODY))
        for b in s.get('bullets', []):
            out.append(Paragraph('\u2022&nbsp;&nbsp;' + mu(b), BULLET))
        if s.get('closing'):
            out.append(Paragraph(mu(s['closing']), BODY))
        out.append(HRFlowable(width='100%', thickness=0.6, color=RULE, spaceBefore=4, spaceAfter=6))
        return out

    def name_cell(row):
        if row.get('parts'):
            kids = []
            for i, p in enumerate(row['parts']):
                if i: kids.append(' / ')
                kids.append("<a href='%s' color='%s'><u>%s</u></a>" % (p['url'], LINK, esc(p['label'])))
            return ''.join(kids)
        if row.get('url'):
            return "<a href='%s' color='%s'><u>%s</u></a>" % (row['url'], LINK, esc(row['name']))
        return esc(row.get('name', ''))

    def sec_policy_table():
        rows = cfg.get('policy_rows', [])
        out = [Paragraph(labels.get('policy_table', 'Policy activities'), H1)]
        leg = cfg.get('scoring', {}).get('legend')
        if leg:
            out.append(Paragraph(mu(leg), BODY_SM))
        out.append(Spacer(1, 3))
        head = [Paragraph('Ref', TH), Paragraph('Policy activity', TH),
                Paragraph('Relevance', TH_C), Paragraph(status_label, TH_C),
                Paragraph(labels.get('why_column', 'Why / how it is (or could be) informed by MO research'), TH)]
        data = [head]
        for r in rows:
            b = band_of(r['relevance'])
            chip = Paragraph('%s<br/><font size=6>%s</font>' % (r['relevance'],
                             {'H': 'HIGH', 'M': 'MED', 'L': 'LOW'}[b]), CHIP)
            stp = Paragraph(esc(r['status']), CHIP)
            data.append([Paragraph(esc(r.get('ref', '')), C_REF), Paragraph(name_cell(r), C_NAME),
                         chip, stp, Paragraph(mu(r.get('why', '')), C_BODY)])
        cw = [48, 166, 58, 72, USABLE - (48 + 166 + 58 + 72)]
        tbl = Table(data, colWidths=cw, repeatRows=1)
        ts = [('BACKGROUND', (0, 0), (-1, 0), TEAL), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
              ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('VALIGN', (2, 1), (3, -1), 'MIDDLE'),
              ('LEFTPADDING', (0, 0), (-1, -1), 5), ('RIGHTPADDING', (0, 0), (-1, -1), 5),
              ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
              ('TOPPADDING', (0, 0), (-1, 0), 6), ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
              ('LINEBELOW', (0, 0), (-1, 0), 0.8, TEAL_DK),
              ('LINEBELOW', (0, 1), (-1, -1), 0.4, colors.white),
              ('GRID', (0, 0), (-1, -1), 0, colors.white)]
        for i, r in enumerate(rows, start=1):
            b = band_of(r['relevance'])
            ts.append(('BACKGROUND', (0, i), (-1, i), BAND_BG[b]))
            ts.append(('BACKGROUND', (2, i), (2, i), BAND_CHIP[b]))
            ts.append(('BACKGROUND', (3, i), (3, i), STATUS_BG.get(r['status'], C['muted'])))
        tbl.setStyle(TableStyle(ts))
        out.append(tbl)
        return out

    def sec_indicators():
        rows = cfg.get('indicator_rows', [])
        if not rows:
            return []
        blk = [Spacer(1, 9), Paragraph(labels.get('indicators', 'Policy indicators'), H1)]
        if cfg.get('indicators_intro'):
            blk.append(Paragraph(mu(cfg['indicators_intro']), BODY_SM))
        blk.append(Spacer(1, 3))
        head = [Paragraph('Indicator (ref)', IND_TH), Paragraph('Indicator name', IND_TH),
                Paragraph('Source framework', IND_TH),
                Paragraph(labels.get('indicator_app_column', 'Current policy application(s)'), IND_TH)]
        data = [head]
        for r in rows:
            data.append([Paragraph(esc(r.get('ref', '')), IND_REF), Paragraph(mu(r.get('name', '')), IND_BODY),
                         Paragraph(mu(r.get('framework', '')), IND_BODY),
                         Paragraph(mu(r.get('application', '')), IND_BODY)])
        cw = [62, 268, 188, USABLE - (62 + 268 + 188)]
        tbl = Table(data, colWidths=cw, repeatRows=1)
        ts = [('BACKGROUND', (0, 0), (-1, 0), TEAL), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
              ('LEFTPADDING', (0, 0), (-1, -1), 5), ('RIGHTPADDING', (0, 0), (-1, -1), 5),
              ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
              ('TOPPADDING', (0, 0), (-1, 0), 6), ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
              ('LINEBELOW', (0, 0), (-1, 0), 0.8, TEAL_DK),
              ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, C['alt_tint']]),
              ('LINEBELOW', (0, 1), (-1, -1), 0.3, RULE)]
        tbl.setStyle(TableStyle(ts)); blk.append(tbl)
        return [KeepTogether(blk)]

    def sec_methods():
        out = [Spacer(1, 9), Paragraph(labels.get('methods', 'Methods'), H1)]
        for m in cfg.get('methods', []):
            out.append(Paragraph(mu(m), BODY_SM))
        if cfg.get('gaps'):
            out.append(Spacer(1, 4))
            out.append(Paragraph(labels.get('gaps', 'Main gaps and caveats'), SUBH))
            out.append(Paragraph(mu(cfg['gaps']), BODY_SM))
        return out

    def sec_references():
        out = [Spacer(1, 8), Paragraph(labels.get('references', 'Additional references used'), H1)]
        if cfg.get('references_primary'):
            out.append(Paragraph(labels.get('references_primary', 'Primary references'),
                                 st('rh1', parent=BODY, fontName='Helvetica-Bold', textColor=TEAL_DK,
                                    fontSize=9.3, spaceAfter=3)))
            for i, r in enumerate(cfg['references_primary'], 1):
                out.append(Paragraph('%d.&nbsp;&nbsp;%s' % (i, mu(r)), REFP))
        if cfg.get('references_other'):
            out.append(Spacer(1, 3))
            out.append(Paragraph(labels.get('references_other', 'Other sources'),
                                 st('rh2', parent=BODY, fontName='Helvetica-Bold', textColor=TEAL_DK,
                                    fontSize=9.3, spaceAfter=3)))
            for i, r in enumerate(cfg['references_other'], 1):
                out.append(Paragraph('%d.&nbsp;&nbsp;%s' % (i, mu(r)), REFP))
        return out

    builders = dict(key_points=sec_key_points, summary=sec_summary, policy_table=sec_policy_table,
                    indicators=sec_indicators, methods=sec_methods, references=sec_references)
    story = [Spacer(1, 2)]
    for name in cfg['section_order']:
        if name in builders:
            story += builders[name]()
    doc.build(story)
    return out_path

# =============================================================================  DOCX
def build_docx(cfg, out_path):
    from docx import Document
    from docx.shared import Pt, RGBColor, Twips
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    T = cfg['theme']; FONT = T.get('font', 'Arial')
    BAND_BG = {'H': T['high_bg'], 'M': T['med_bg'], 'L': T['low_bg']}
    BAND_CHIP = {'H': T['high_chip'], 'M': T['med_chip'], 'L': T['low_chip']}
    _sa = cfg.get('status_axis') or {}
    status_label = _sa.get('label', 'Status')
    STATUS_BG = _sa.get('values') or {'Current': T['status_current'], 'Opportunity': T['status_opportunity']}
    labels = cfg.get('labels', {}); meta = cfg.get('meta', {})

    def setfont(run, size, color=None, bold=False, italic=False):
        run.font.name = FONT; run.font.size = Pt(size)
        r = run._element.rPr.rFonts if run._element.rPr is not None and run._element.rPr.rFonts is not None else None
        rpr = run._element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
            rf.set(qn(a), FONT)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        run.bold = bold; run.italic = italic

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill)
        tcPr.append(shd)

    def set_cell_margins(cell, top=60, bottom=60, left=110, right=110):
        tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
        for tag, val in (('top', top), ('bottom', bottom), ('start', left), ('end', right),
                         ('left', left), ('right', right)):
            e = OxmlElement('w:' + tag); e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa'); m.append(e)
        tcPr.append(m)

    def no_borders_table(tbl, color='FFFFFF', size=4):
        tblPr = tbl._tbl.tblPr; b = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), str(size)); e.set(qn('w:color'), color); b.append(e)
        tblPr.append(b)

    def header_row(row):
        trPr = row._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader')
        h.set(qn('w:val'), 'true'); trPr.append(h)

    def add_runs(paragraph, text, size, base_color=None):
        for sp in parse_markup(text):
            if sp['url']:
                add_hyperlink(paragraph, sp['url'], sp['text'], size, sp['bold'], sp['italic'])
            else:
                run = paragraph.add_run(sp['text'])
                setfont(run, size, base_color, sp['bold'], sp['italic'])

    def add_hyperlink(paragraph, url, text, size, bold, italic):
        part = paragraph.part
        r_id = part.relate_to(url,
            'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
            is_external=True)
        hl = OxmlElement('w:hyperlink'); hl.set(qn('r:id'), r_id)
        r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
        rf = OxmlElement('w:rFonts')
        for a in ('w:ascii', 'w:hAnsi', 'w:cs'): rf.set(qn(a), FONT)
        rPr.append(rf)
        col = OxmlElement('w:color'); col.set(qn('w:val'), T['link']); rPr.append(col)
        u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
        if bold: rPr.append(OxmlElement('w:b'))
        if italic: rPr.append(OxmlElement('w:i'))
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); rPr.append(sz)
        r.append(rPr)
        t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text; r.append(t)
        hl.append(r); paragraph._p.append(hl)

    def para(container, text=None, size=9, color=None, bold=False, italic=False,
             align=None, space_after=4, space_before=0):
        p = container.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(space_before)
        if align: p.alignment = align
        if text is not None:
            if any(tok in text for tok in ('**', '*', '[[')):
                add_runs(p, text, size, color)
                if bold:
                    for r in p.runs: r.bold = True
            else:
                run = p.add_run(text); setfont(run, size, color, bold, italic)
        return p

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Twips(15840), Twips(12240)
    sec.left_margin = Twips(720); sec.right_margin = Twips(720)
    sec.top_margin = Twips(1080); sec.bottom_margin = Twips(900)
    sec.header_distance = Twips(540); sec.footer_distance = Twips(360)
    CONTENT = 15840 - 720 - 720  # 14400 dxa

    style = doc.styles['Normal']; style.font.name = FONT; style.font.size = Pt(9)
    style.font.color.rgb = RGBColor.from_string(T['ink'])

    # running header / footer
    hp = sec.header.paragraphs[0]; hp.paragraph_format.space_after = Pt(0)
    run = hp.add_run(meta.get('title', '')); setfont(run, 7.5, T['muted'])
    pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6'); bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), T['teal']); pbdr.append(bottom)
    hp._p.get_or_add_pPr().append(pbdr)

    fp = sec.footer.paragraphs[0]; fp.paragraph_format.space_before = Pt(0)
    tabs = fp.paragraph_format.tab_stops
    from docx.enum.text import WD_TAB_ALIGNMENT
    tabs.add_tab_stop(Twips(CONTENT), WD_TAB_ALIGNMENT.RIGHT)
    r1 = fp.add_run(meta.get('footer', '')); setfont(r1, 7, T['muted'])
    r2 = fp.add_run('\tPage '); setfont(r2, 7, T['muted'])
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); fp._p.append(fld)
    fpbdr = OxmlElement('w:pBdr'); top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single'); top.set(qn('w:sz'), '4'); top.set(qn('w:space'), '4')
    top.set(qn('w:color'), T['rule']); fpbdr.append(top); fp._p.get_or_add_pPr().append(fpbdr)

    def H1(text):
        para(doc, text, size=13, color=T['teal_dk'], bold=True, space_before=11, space_after=6)

    def one_cell_table(fill, builder, left_accent=None):
        tbl = doc.add_table(rows=1, cols=1); tbl.autofit = False
        tbl.columns[0].width = Twips(CONTENT)
        cell = tbl.cell(0, 0); cell.width = Twips(CONTENT)
        shade(cell, fill); set_cell_margins(cell, 140, 120, 200, 200)
        # borders
        tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
        for edge in ('top', 'bottom', 'right'):
            e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4')
            e.set(qn('w:color'), T['teal']); b.append(e)
        le = OxmlElement('w:left'); le.set(qn('w:val'), 'single')
        le.set(qn('w:sz'), str(24 if left_accent else 4)); le.set(qn('w:color'), left_accent or T['teal'])
        b.append(le); tcPr.append(b)
        cell._element.clear_content() if False else None
        # remove default empty paragraph
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
        builder(cell)

    # ---- section builders -------------------------------------------------
    def sec_title_band():
        tbl = doc.add_table(rows=1, cols=1); tbl.autofit = False
        tbl.columns[0].width = Twips(CONTENT); cell = tbl.cell(0, 0); cell.width = Twips(CONTENT)
        shade(cell, T['teal']); set_cell_margins(cell, 160, 160, 200, 200)
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
        para(cell, meta.get('title', ''), size=15, color='FFFFFF', bold=True, space_after=2)
        sub = (meta.get('subtitle', '') + ('  \u00b7  ' + meta.get('prepared', '') if meta.get('prepared') else ''))
        para(cell, sub, size=9, color='D6E6E9', space_after=0)
        para(doc, '', space_after=6)

    def sec_key_points():
        items = cfg.get('key_points', [])
        if not items: return
        def build(cell):
            para(cell, labels.get('key_points', 'Key points'), size=11, color=T['teal_dk'],
                 bold=True, space_after=3)
            for it in items:
                p = cell.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.left_indent = Twips(220); p.paragraph_format.first_line_indent = Twips(-150)
                add_runs(p, it, 9)
        one_cell_table(T['kp_fill'], build, left_accent=T['teal'])
        para(doc, '', space_after=4)

    def sec_summary():
        s = cfg.get('summary', {})
        H1(labels.get('summary', 'Summary'))
        if s.get('intro'): para(doc, s['intro'], size=8.5, space_after=6)
        for b in s.get('bullets', []):
            p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(4)
            add_runs(p, b, 8.5)
        if s.get('closing'): para(doc, s['closing'], size=8.5, space_after=6)

    def cell_para(cell, text=None, size=8.3, color=None, bold=False, align=None, chip=False, big=None, small=None):
        if cell.paragraphs and not cell.paragraphs[0].runs and chip is False and text is not None and len(cell.paragraphs) == 1 and cell.paragraphs[0].text == '':
            cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
        p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        if align: p.alignment = align
        if chip:
            r = p.add_run(big); setfont(r, 10, color, True)
            if small:
                r2 = p.add_run(); r2.add_break(); setfont(r2, 6, color, True)
                r3 = p.add_run(small); setfont(r3, 6, color, True)
            return p
        if text is not None:
            if any(tok in text for tok in ('**', '*', '[[')):
                add_runs(p, text, size, color)
                if bold:
                    for r in p.runs: r.bold = True
            else:
                run = p.add_run(text); setfont(run, size, color, bold)
        return p

    def fresh_cell(cell):
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)

    def sec_policy_table():
        rows = cfg.get('policy_rows', [])
        H1(labels.get('policy_table', 'Policy activities'))
        leg = cfg.get('scoring', {}).get('legend')
        if leg: para(doc, leg, size=8.5, space_after=6)
        widths = [900, 3100, 1300, 1400, CONTENT - (900 + 3100 + 1300 + 1400)]
        tbl = doc.add_table(rows=1, cols=5); tbl.autofit = False; no_borders_table(tbl)
        hdr = tbl.rows[0]; header_row(hdr)
        htexts = ['Ref', 'Policy activity', 'Relevance', status_label,
                  labels.get('why_column', 'Why / how it is (or could be) informed by MO research')]
        for i, c in enumerate(hdr.cells):
            c.width = Twips(widths[i]); shade(c, T['teal']); set_cell_margins(c)
            fresh_cell(c)
            cell_para(c, htexts[i], size=8.6, color='FFFFFF', bold=True,
                      align=WD_ALIGN_PARAGRAPH.CENTER if i in (2, 3) else None)
        for r in rows:
            b = band_of(r['relevance']); tr = tbl.add_row()
            for i, c in enumerate(tr.cells):
                c.width = Twips(widths[i]); set_cell_margins(c); fresh_cell(c)
            c0, c1, c2, c3, c4 = tr.cells
            shade(c0, BAND_BG[b]); cell_para(c0, r.get('ref', ''), bold=True)
            shade(c1, BAND_BG[b])
            if r.get('parts'):
                p = c1.add_paragraph(); p.paragraph_format.space_after = Pt(0)
                for j, pt in enumerate(r['parts']):
                    if j: setfont(p.add_run(' / '), 8.3, T['ink'])
                    add_hyperlink(p, pt['url'], pt['label'], 8.3, True, False)
            else:
                txt = '[[%s|%s]]' % (r['name'], r['url']) if r.get('url') else r['name']
                cp = cell_para(c1, txt, bold=not r.get('url'))
            c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            c3.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            shade(c2, BAND_CHIP[b])
            cell_para(c2, chip=True, big=str(r['relevance']), small={'H': 'HIGH', 'M': 'MED', 'L': 'LOW'}[b],
                      color='FFFFFF', align=WD_ALIGN_PARAGRAPH.CENTER)
            shade(c3, STATUS_BG.get(r['status'], T['muted']))
            sp = cell_para(c3, align=WD_ALIGN_PARAGRAPH.CENTER); rr = sp.add_run(r['status'])
            setfont(rr, 7.5, 'FFFFFF', True)
            shade(c4, BAND_BG[b]); cell_para(c4, r.get('why', ''))
        for i, w in enumerate(widths):
            for cell in tbl.columns[i].cells: cell.width = Twips(w)

    def sec_indicators():
        rows = cfg.get('indicator_rows', [])
        if not rows: return
        H1(labels.get('indicators', 'Policy indicators'))
        if cfg.get('indicators_intro'): para(doc, cfg['indicators_intro'], size=8.5, space_after=6)
        widths = [1200, 5000, 3500, CONTENT - (1200 + 5000 + 3500)]
        tbl = doc.add_table(rows=1, cols=4); tbl.autofit = False; no_borders_table(tbl, color=T['rule'], size=2)
        hdr = tbl.rows[0]; header_row(hdr)
        htexts = ['Indicator (ref)', 'Indicator name', 'Source framework',
                  labels.get('indicator_app_column', 'Current policy application(s)')]
        for i, c in enumerate(hdr.cells):
            c.width = Twips(widths[i]); shade(c, T['teal']); set_cell_margins(c); fresh_cell(c)
            cell_para(c, htexts[i], size=8.6, color='FFFFFF', bold=True)
        for k, r in enumerate(rows):
            tr = tbl.add_row(); fill = T['alt_tint'] if k % 2 else 'FFFFFF'
            vals = [r.get('ref', ''), r.get('name', ''), r.get('framework', ''), r.get('application', '')]
            for i, c in enumerate(tr.cells):
                c.width = Twips(widths[i]); set_cell_margins(c); shade(c, fill); fresh_cell(c)
                cell_para(c, vals[i], bold=(i == 0))
        for i, w in enumerate(widths):
            for cell in tbl.columns[i].cells: cell.width = Twips(w)

    def sec_methods():
        H1(labels.get('methods', 'Methods'))
        for m in cfg.get('methods', []):
            p = doc.add_paragraph(style='List Number'); p.paragraph_format.space_after = Pt(4)
            add_runs(p, m, 8.5)
        if cfg.get('gaps'):
            para(doc, labels.get('gaps', 'Main gaps and caveats'), size=9, color=T['teal_dk'],
                 bold=True, space_before=4, space_after=3)
            para(doc, cfg['gaps'], size=8.5, space_after=6)

    def sec_references():
        H1(labels.get('references', 'Additional references used'))
        if cfg.get('references_primary'):
            para(doc, labels.get('references_primary', 'Primary references'), size=9, color=T['teal_dk'],
                 bold=True, space_after=3)
            for r in cfg['references_primary']:
                p = doc.add_paragraph(style='List Number'); p.paragraph_format.space_after = Pt(3)
                add_runs(p, r, 8)
        if cfg.get('references_other'):
            para(doc, labels.get('references_other', 'Other sources'), size=9, color=T['teal_dk'],
                 bold=True, space_before=6, space_after=3)
            for r in cfg['references_other']:
                p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3)
                add_runs(p, r, 8)

    builders = dict(key_points=sec_key_points, summary=sec_summary, policy_table=sec_policy_table,
                    indicators=sec_indicators, methods=sec_methods, references=sec_references)
    sec_title_band()
    for name in cfg['section_order']:
        if name in builders:
            builders[name]()
    doc.save(out_path)
    return out_path

# =============================================================================  main
def main():
    ap = argparse.ArgumentParser(description='Render assessment to PDF and/or Word from YAML.')
    ap.add_argument('config')
    ap.add_argument('--format', choices=['pdf', 'docx', 'both'], default='both')
    ap.add_argument('--outdir', default='.')
    args = ap.parse_args()
    cfg = load_config(args.config)
    base = cfg.get('meta', {}).get('output_basename', 'assessment')
    os.makedirs(args.outdir, exist_ok=True)
    made = []
    if args.format in ('pdf', 'both'):
        made.append(build_pdf(cfg, os.path.join(args.outdir, base + '.pdf')))
    if args.format in ('docx', 'both'):
        made.append(build_docx(cfg, os.path.join(args.outdir, base + '.docx')))
    for m in made:
        print('wrote', m)

if __name__ == '__main__':
    main()
