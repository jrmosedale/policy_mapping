#!/usr/bin/env python3
"""
verify_toolkit.py — prove this toolkit works on this machine, in one command.

    python3 PA_toolkit/code/verify_toolkit.py

Checks dependencies, confirms render.py is intact, renders both worked-example
templates to PDF and DOCX in a temporary directory, and inspects the results.
Writes nothing into the project. Exit 0 = the toolkit is healthy.

WHY THIS EXISTS
---------------
render.py is ~32 KB / ~600 lines. Some AI assistants truncate large attachments
without saying so, see only the first screenful, and report that the file is
"only a scaffold" with build_pdf/build_docx unimplemented. That claim is wrong,
and this script is the two-second refutation: if it renders real documents, the
implementation is present.

If an assistant tells you render.py is incomplete, run this first. If it passes,
the assistant did not read the whole file. Give it the file again in chunks, or
paste the specific function it says is missing, rather than letting it rewrite a
working renderer -- a rewrite silently changes the layout and scoring presentation
that make separate assessments comparable, which is the whole point of the kit.
"""
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent          # PA_toolkit/code
KIT = HERE.parent                                # PA_toolkit
RENDER = HERE / "render.py"
TEMPLATES = [KIT / "template_files" / "TEMPLATE_applied.yaml",
             KIT / "template_files" / "TEMPLATE_gap.yaml"]

# Landmarks that must be present in a complete render.py.
REQUIRED = ["def parse_markup(", "def band_of(", "DEFAULT_THEME", "def load_config(",
            "def build_pdf(", "def build_docx(", "def main("]
MIN_LINES = 500

fails = []


def check(label, ok, detail=""):
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}" + (f" -- {detail}" if detail else ""))
    if not ok:
        fails.append(label)
    return ok


def main():
    print("\n== 1. Dependencies ==")
    for mod, pkg in [("reportlab", "reportlab>=4.0"),
                     ("docx", "python-docx>=1.1"),
                     ("yaml", "PyYAML>=6.0")]:
        try:
            __import__(mod)
            check(pkg, True)
        except ImportError:
            check(pkg, False, "not installed -- pip install -r PA_toolkit/code/requirements.txt")

    print("\n== 2. render.py is intact ==")
    if not check("render.py present", RENDER.exists(), str(RENDER)):
        return report()
    src = RENDER.read_text(encoding="utf-8")
    n = len(src.splitlines())
    check(f"length ({n} lines, {len(src.encode())} bytes)", n >= MIN_LINES,
          f"expected >= {MIN_LINES}; a short file means a truncated copy")
    for token in REQUIRED:
        check(f"contains {token!r}", token in src)
    for stub in ("NotImplementedError", "TODO:", "FIXME:"):
        check(f"no {stub}", stub not in src)

    print("\n== 3. Render both templates ==")
    with tempfile.TemporaryDirectory() as tmp:
        for tpl in TEMPLATES:
            if not check(f"{tpl.name} present", tpl.exists()):
                continue
            r = subprocess.run([sys.executable, str(RENDER), str(tpl),
                                "--format", "both", "--outdir", tmp],
                               capture_output=True, text=True)
            if not check(f"{tpl.name} renders", r.returncode == 0,
                         (r.stderr.strip().splitlines() or [""])[-1]):
                continue
            made = sorted(Path(tmp).glob("*"))
            pdfs = [p for p in made if p.suffix == ".pdf"]
            docs = [p for p in made if p.suffix == ".docx"]
            check(f"{tpl.name} -> PDF", bool(pdfs) and pdfs[-1].stat().st_size > 5000,
                  f"{pdfs[-1].name} {pdfs[-1].stat().st_size // 1024} KB" if pdfs else "none")
            check(f"{tpl.name} -> DOCX", bool(docs) and docs[-1].stat().st_size > 5000,
                  f"{docs[-1].name} {docs[-1].stat().st_size // 1024} KB" if docs else "none")
            if docs:
                try:
                    from docx import Document
                    d = Document(str(docs[-1]))
                    check(f"{tpl.name} DOCX has tables", len(d.tables) >= 2,
                          f"{len(d.tables)} tables, {len(d.paragraphs)} paragraphs")
                except Exception as e:
                    check(f"{tpl.name} DOCX opens", False, str(e)[:60])
            for p in made:
                p.unlink()
    return report()


def report():
    print()
    if fails:
        print(f"FAILED -- {len(fails)} check(s): " + "; ".join(fails))
        print("\nIf the failures are dependencies, install them and re-run.")
        print("If render.py itself is short or missing landmarks, your copy is truncated:")
        print("re-copy it from the original project rather than reconstructing it.")
        return 1
    print("ALL CHECKS PASSED -- render.py is complete and the toolkit renders on this machine.")
    print("\nIf an AI assistant claims render.py is only a scaffold, it has not read the whole")
    print("file. Do not let it rewrite the renderer: re-supply the file instead.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
