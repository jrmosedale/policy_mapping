#!/usr/bin/env python3
"""Apply proposal_data.py to copies of the governance workbooks.

  python3 apply_proposal.py SRC_DIR OUT_DIR

Reads the canonical files from SRC_DIR and writes edited copies to OUT_DIR (same filenames).
Follows WORKBOOK_WRITE_PROTOCOL §4: append rows only (no insert_rows), stripe by row parity,
copy style from a same-parity template row, never bracket key columns. Bumps Changelog!B2 and
appends a Changelog row. Archiving and finalise.py are done by the caller.
"""
import sys, copy, datetime
from pathlib import Path
import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parent))
import proposal_data as d

UK = "uk_climate_nature_governance.xlsx"
INTL = "international_climate_nature_governance.xlsx"
TODAY = datetime.date.today().isoformat()

def style_from(dst, src):
    dst.font = copy.copy(src.font); dst.fill = copy.copy(src.fill); dst.border = copy.copy(src.border)
    dst.alignment = copy.copy(src.alignment); dst.number_format = src.number_format

def headers(ws):
    return [c.value for c in ws[1]]

def append_records(ws, recs):
    h = headers(ws)
    existing = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
    for rec in recs:
        assert rec["Record_ID"] not in existing, f"{rec['Record_ID']} already exists"
        unknown = [k for k in rec if not k.startswith("_") and k not in h]
        assert not unknown, f"{rec['Record_ID']}: unknown columns {unknown}"
        r = ws.max_row + 1
        tmpl = 2 if r % 2 == 0 else 3
        for ci, col in enumerate(h, start=1):
            c = ws.cell(r, ci)
            style_from(c, ws.cell(tmpl, ci))
            v = rec.get(col)
            c.value = v if v not in ("", None) else None

def find_row(ws, rid):
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == rid:
            return r
    raise KeyError(rid)

def apply_edits(wb, edits):
    for _, sheet, rid, col, old, new, _why in edits:
        ws = wb[sheet]; h = headers(ws); r = find_row(ws, rid); c = ws.cell(r, h.index(col) + 1)
        if old == "append":
            c.value = (c.value or "") + new
        else:
            if old is not None:
                assert old in str(c.value or ""), f"{rid}.{col}: expected '{old}', found '{c.value}'"
            c.value = new

def bump(wb, summary):
    cl = wb["Changelog"]
    old = cl["B2"].value
    new = int(old) + 1
    cl["B2"].value = str(new) if isinstance(old, str) else new
    for r in range(1, 6):
        if cl.cell(r, 1).value == "Last updated":
            cl.cell(r, 2).value = TODAY
    last = cl.max_row; nr = last + 1
    for ci in range(1, 5):
        style_from(cl.cell(nr, ci), cl.cell(last, ci))
    cl.cell(nr, 1).value = f"v{new}"; cl.cell(nr, 2).value = TODAY
    cl.cell(nr, 3).value = f"v{new}: {summary}"; cl.cell(nr, 4).value = "Cowork (Claude)"
    return int(old), new

UK_SUMMARY = ("Open work item 1 — approved queue candidates written. NEW: 4 Legislation (LEG-068 Natural Environment (Scotland) Act 2026; "
 "LEG-069 Environment (Principles, Governance and Biodiversity Targets) (Wales) Act 2026; LEG-070 Building Act 1984, parent of Part O; "
 "LEG-071 NI Climate Commissioner Regulations 2025, now ORG-031's Statutory_Role), 5 Public Bodies (ORG-045 Environmental Standards Scotland; "
 "ORG-046 Office of Environmental Governance Wales; ORG-047 UKHSA; ORG-048 Local Nature Partnerships (collective); ORG-049 Department of "
 "Health and Social Care, UKHSA's sponsor), 39 Policies & Activities (POL-072–POL-110: weather-health alerts, Adverse Weather and Health Plan, Part O, NSWWS, "
 "UK Plant Health Risk Register, plant health contingency plans, 30by30 England and its assessment guidance, nature security assessment, Farming "
 "Roadmap 2050, Water White Paper 2026, Environment Act target delivery plans (overview + 13), GENP, deforestation regulations approach, Biodiverse "
 "Landscapes Fund, TE2100, Blue Belt, Peatland Restoration Sector Capacity Grant, NFWR 2025, Species Recovery Programme, Local Sites, NI Peatland "
 "Strategy 2040, GGR review and response, Scottish Biodiversity Strategy 2045 and Delivery Plan 2024–30, 30 by 30 Scotland). Policies & Activities "
 "General_Type gains 'Operational service / System' (approved 10 Sep 2026). CORRECTION: LEG-034 described a non-existent 'Environment (Scotland) "
 "Act 2023' whose link resolved to the Bail and Release from Custody (Scotland) Act 2023; its provisions (ESS, environmental principles) are those of "
 "the UK Withdrawal from the European Union (Continuity) (Scotland) Act 2021, so the ID was repurposed to that Act and its inbound references kept; "
 "[LEG-034] removed from ORG-014 Statutory_Role. Approved by Jonathan 6 Oct 2026 (Scottish Biodiversity Programme rejected; JNCC NI peatland review replaced by POL-106; CCC-ESS MoU recorded in ORG-045 prose; NI Nature Recovery Strategy held pending adoption). Every value verified at source 6 Oct 2026; proposal and verification record in "
 "outputs_other/workbook_proposal_2026-10-06/.")
INTL_SUMMARY = ("INT-L-007 UK_Links gains [POL-078] 30by30 on land in England and [POL-110] 30 by 30 in Scotland, reciprocal to their new Intl_Links "
 "(UK governance v27).")

def main(src, out):
    src, out = Path(src), Path(out); out.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(src / UK)
    append_records(wb["Legislation"], d.LEG)
    append_records(wb["Public Bodies"], d.ORG)
    append_records(wb["Policies & Activities"], d.POL)
    apply_edits(wb, d.EDITS)
    dd = wb["Data Dictionary"]
    for _, sheet, col, term, _why in d.VOCAB:
        for r in dd.iter_rows():
            if r[0].value == sheet and r[1].value == col:
                assert term not in r[4].value; r[4].value = r[4].value + " | " + term
    ov, nv = bump(wb, UK_SUMMARY)
    wb.save(out / UK); print(f"UK v{ov} -> v{nv}")
    wi = openpyxl.load_workbook(src / INTL)
    apply_edits(wi, d.EDITS_INTL)
    ov, nv = bump(wi, INTL_SUMMARY)
    wi.save(out / INTL); print(f"Intl v{ov} -> v{nv}")

if __name__ == "__main__":
    main(*sys.argv[1:3])
