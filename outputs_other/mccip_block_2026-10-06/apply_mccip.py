#!/usr/bin/env python3
"""Write the MCCIP block into copies of the canonical workbooks: python3 apply_mccip.py SRC_DIR OUT_DIR"""
import sys, copy, datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill
sys.path.insert(0, str(Path(__file__).resolve().parent)); import mccip_data as d
IND="indicators_climate_nature.xlsx"; UK="uk_climate_nature_governance.xlsx"; TODAY=datetime.date.today().isoformat()
def sty(dst,src): dst._style=copy.copy(src._style)
def bump(wb,summary):
    cl=wb["Changelog"]; old=cl["B2"].value; new=int(old)+1; cl["B2"].value=str(new) if isinstance(old,str) else new
    for r in range(1,6):
        if cl.cell(r,1).value=="Last updated": cl.cell(r,2).value=TODAY
    last=cl.max_row; nr=last+1
    for ci in range(1,5): sty(cl.cell(nr,ci),cl.cell(last,ci))
    cl.cell(nr,1).value=f"v{new}"; cl.cell(nr,2).value=TODAY; cl.cell(nr,3).value=f"v{new}: {summary}"; cl.cell(nr,4).value="Cowork (Claude)"
    return old,new
def main(src,out):
    src,out=Path(src),Path(out); out.mkdir(parents=True,exist_ok=True)
    wb=openpyxl.load_workbook(src/IND)
    assert "MCCIP Indicators" not in wb.sheetnames
    tmpl=wb["EEA Biodiversity Indicators"]
    ws=wb.copy_worksheet(tmpl); ws.title="MCCIP Indicators"
    # move to just after SDG Indicators
    wb.move_sheet(ws, offset=wb.sheetnames.index("SDG Indicators")+1-wb.sheetnames.index("MCCIP Indicators"))
    ws.delete_rows(2, ws.max_row)          # new sheet: remove copied data rows (nothing below to shift)
    nc=len(d.COLS)
    if ws.max_column>nc: ws.delete_cols(nc+1, ws.max_column-nc)
    head=PatternFill("solid", fgColor="FF0B5563")
    for ci,col in enumerate(d.COLS,1):
        c=ws.cell(1,ci); sty(c,tmpl.cell(1,min(ci,tmpl.max_column))); c.value=col; c.fill=head
    widths={"Record_ID":12,"Indicator_Name":32,"Source_Framework":30,"NCF_Category":16,"Policy_Sector":30,"Indicator_Type":22,"Units_Measure":34,
            "Scope":12,"Policy_Goal":45,"Source_Reference":40,"Data_Source":35,"Climate_Score":9,"Climate_Dependency":14,"Climate_Rationale":70,
            "Policy_Links":26,"Framework_Classification":55,"Source_URL":40}
    for ci,col in enumerate(d.COLS,1): ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width=widths[col]
    for i,rec in enumerate(d.ROWS):
        r=i+2; par=2 if r%2==0 else 3
        for ci,col in enumerate(d.COLS,1):
            c=ws.cell(r,ci); sty(c,tmpl.cell(par,min(ci,tmpl.max_column))); c.value=rec[col]
        ws.row_dimensions[r].height=90
    ws.auto_filter.ref=f"A1:{openpyxl.utils.get_column_letter(nc)}1"
    # register
    rg=wb["Indicator Framework"]; h=[c.value for c in rg[1]]; r=rg.max_row+1; par=2 if r%2==0 else 3
    for ci,col in enumerate(h,1): sty(rg.cell(r,ci),rg.cell(par,ci)); rg.cell(r,ci).value=d.REGISTER.get(col)
    # data dictionary
    dd=wb["Data Dictionary"]
    def ddset(sheet,col,fn_sheet=None,fn_vocab=None):
        for row in dd.iter_rows():
            if row[0].value==sheet and row[1].value==col:
                if fn_sheet: row[0].value=fn_sheet(row[0].value)
                if fn_vocab: row[4].value=fn_vocab(row[4].value)
                return
        raise KeyError((sheet,col))
    ddset("All framework sheets","Record_ID",fn_vocab=lambda v: v+" | IND-M-NNN (MCCIP)")
    ddset("EIF / CCC / JNCC / SoN / BIP","Framework_Classification",fn_sheet=lambda v: v+" / MCCIP")
    ddset("CCC / JNCC","Source_URL",fn_sheet=lambda v: v+" / MCCIP")
    ddset("Indicator Framework","Framework_ID",fn_vocab=lambda v: v.replace("IFW-01…IFW-09","IFW-01…IFW-12"))
    o,n=bump(wb,"MCCIP indicator block added (approved by Jonathan 6 Oct 2026: the 16 physical-environment and ecosystem-change topics). "
      "New register row [IFW-12] 'MCCIP UK Marine Climate Change Impacts Evidence Hub', new sheet 'MCCIP Indicators' (placed after SDG Indicators), new ID family IND-M "
      "(IND-M-001–010 physical environment, IND-M-011–016 ecosystem change). Assessment-derived like [IFW-07]: each record is a standing MCCIP topic review; "
      "Framework_Classification keeps MCCIP's separate observed and projection confidence ratings; Climate_Rationale gives the Legend rule applied and the observed and projected statements. "
      "Climate scores checked against outputs_other/CLIMATE_SCORE_METHOD.md: 3 for the nine physical climate-system variables (rule E ii) and for seabirds and waterbirds (rule F: the wintering-waterbird component governs, aligned with IND-J-013); 2 for plankton, fish and marine mammals (rule D); 1 for coastal geomorphology, coastal and intertidal habitats, and shallow/shelf/deep-sea habitats (rule B). Rule G siblings cited in each rationale. "
      "No new controlled-vocabulary terms; Indicator_Type uses existing terms. The six societal-impact topics are excluded (level-of-generality rule). GBF_Targets_Clean not carried (0/16). "
      "Data Dictionary: IND-M family declared; Framework_Classification and Source_URL rows extended to MCCIP; Framework_ID range updated to IFW-12. Every value read from the MCCIP topic pages on 6 Oct 2026.")
    wb.save(out/IND); print("indicators",o,"->",n)
    wu=openpyxl.load_workbook(src/UK); pa=wu["Policies & Activities"]; hh=[c.value for c in pa[1]]
    for r in range(2,pa.max_row+1):
        if pa.cell(r,1).value=="POL-049":
            c=pa.cell(r,hh.index("Indicator_Frameworks")+1); assert not c.value; c.value="[IFW-12]"
            pa.cell(r,hh.index("Date_Last_Updated")+1).value="October 2026"
    o,n=bump(wu,"POL-049 MCCIP Indicator_Frameworks gains [IFW-12], the MCCIP indicator block added in indicators v28.")
    wu.save(out/UK); print("UK",o,"->",n)
if __name__=="__main__": main(*sys.argv[1:3])
