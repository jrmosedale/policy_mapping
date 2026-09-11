#!/usr/bin/env python3
"""Report [CODE] mentions in narrative prose fields that are NOT captured by any
curated link column consumed by build_governance_diagram.py. Writes Markdown + CSV."""
import openpyxl, re, csv, os
CODE = re.compile(r'\[((?:LEG|ORG|POL|IFW|INT-L|INT-O|INT-P|EU-L)-\d+)\]')

USED = {
 ("UK","Legislation"):["Enabling_Legislation","Related_Legislation","Lead_Department","Intl_Links"],
 ("UK","Public Bodies"):["Statutory_Role","Lead_Department","Indicator_Frameworks","Intl_Links"],
 ("UK","Policies & Activities"):["Enabling_Legislation","Lead_Organisation","Parent_Programme","Funded_By","Related_Policy_Links","Indicator_Frameworks","Intl_Links"],
 ("INTL","Conventions & Treaties"):["Parent_Convention","UK_Links","Intl_Links","Indicator_Frameworks","UK_Ratification"],
 ("INTL","EU Legislation"):["Parent_Convention","UK_Links","Intl_Links","Indicator_Frameworks","UK_Ratification"],
 ("INTL","Associated Bodies"):["Parent_Convention","UK_Links","Intl_Links","Indicator_Frameworks"],
 ("INTL","Intl Policies & Activities"):["Parent_Convention","Lead_Body","UK_Links","Intl_Links","Indicator_Frameworks"],
 ("IND","Indicator Framework"):["Indirect_Policy_Links","Lead_Organisation"],
}
NARRATIVE = ["Key_Provisions","Scope_And_Commitments","Remit_And_Powers","Key_Commitments"]
FILES = {"UK":"uk_climate_nature_governance.xlsx",
         "INTL":"international_climate_nature_governance.xlsx",
         "IND":"indicators_climate_nature.xlsx"}
PK = {"UK":"Record_ID","INTL":"Record_ID","IND":"Framework_ID"}
WBNAME = {"UK":"uk_climate_nature_governance.xlsx",
          "INTL":"international_climate_nature_governance.xlsx",
          "IND":"indicators_climate_nature.xlsx"}

names = {}
rows = []
for wbk, fn in FILES.items():
    wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
    for sh in wb.sheetnames:
        if (wbk, sh) not in USED: continue
        data = list(wb[sh].iter_rows(values_only=True))
        idx = {h:i for i,h in enumerate(data[0]) if h}
        pk = PK[wbk]
        for r in data[1:]:
            rid = r[idx[pk]] if pk in idx else None
            if not rid: continue
            if "Name" in idx and r[idx["Name"]]: names[rid] = r[idx["Name"]]
            if "Source_Framework" in idx and r[idx["Source_Framework"]]: names[rid] = r[idx["Source_Framework"]]
    wb.close()

for wbk, fn in FILES.items():
    wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
    for sh in wb.sheetnames:
        key = (wbk, sh)
        if key not in USED: continue
        data = list(wb[sh].iter_rows(values_only=True))
        idx = {h:i for i,h in enumerate(data[0]) if h}
        pk = PK[wbk]
        narr = [c for c in NARRATIVE if c in idx]
        if not narr: continue
        for r in data[1:]:
            rid = r[idx[pk]] if pk in idx else None
            if not rid: continue
            captured = set()
            for c in USED[key]:
                if c in idx and r[idx[c]]:
                    captured |= set(CODE.findall(str(r[idx[c]])))
            for nc in narr:
                v = r[idx[nc]]
                if not v: continue
                orph = [c for c in dict.fromkeys(CODE.findall(str(v))) if c not in captured]
                if orph:
                    rows.append({
                        "workbook": WBNAME[wbk], "sheet": sh, "record_id": rid,
                        "record_name": names.get(rid,""), "narrative_field": nc,
                        "uncaptured_codes": "; ".join(orph),
                        "n": len(orph),
                        "targets": "; ".join(f"{c} {names.get(c,'')}".strip() for c in orph),
                    })
    wb.close()

with open("orphan_narrative_codes.csv","w",newline="",encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)

# Markdown, grouped by workbook
lines = ["# Uncaptured `[CODE]` mentions in narrative fields","",
         "Codes appearing in narrative prose that are **not** present in any curated link column "
         "consumed by `build_governance_diagram.py` for that record. These are candidate missing links: "
         "each is either (a) a genuine relationship that should be promoted into a link column "
         "(e.g. `Related_Policy_Links`), or (b) an incidental textual mention that should stay prose.","",
         f"**Total: {sum(r['n'] for r in rows)} uncaptured code mentions across {len(rows)} records.**",""]
for wbk in ["UK","INTL","IND"]:
    sub = [r for r in rows if r["workbook"]==WBNAME[wbk]]
    if not sub: continue
    lines += [f"## `{WBNAME[wbk]}`",""]
    for sh in dict.fromkeys(r["sheet"] for r in sub):
        ss = [r for r in sub if r["sheet"]==sh]
        lines += [f"### {sh} — {sum(r['n'] for r in ss)} mentions in {len(ss)} records","",
                  "| Record | Name | Narrative field | Uncaptured codes |","|---|---|---|---|"]
        for r in ss:
            nm = (r["record_name"] or "")[:52]
            lines.append(f"| `{r['record_id']}` | {nm} | `{r['narrative_field']}` | {r['uncaptured_codes']} |")
        lines.append("")
open("orphan_narrative_codes.md","w",encoding="utf-8").write("\n".join(lines))
print(f"{sum(r['n'] for r in rows)} mentions / {len(rows)} records -> orphan_narrative_codes.md + .csv")
