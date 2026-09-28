#!/usr/bin/env python3
"""rebuild_xref.py — regenerate the Cross-Reference Index sheet from its source sheets.

The Cross-Reference Index is a DERIVED sheet: one row per record, carrying a copy of a few
columns from the sheet the record actually lives on. Nothing reads it — neither dashboard
touches it — so when it drifts from the source sheets nothing breaks visibly and nobody
notices. It had drifted badly by September 2026: 141 of 177 UK rows and 87 of 103
international rows carried a Policy_Sector from the pre-migration vocabulary, five UK rows
carried a superseded General_Type, and four records were missing from the index altogether.

Hand-maintaining a derived sheet guarantees that outcome. This script removes the need to.

Rows are written in source order: sheets in workbook order, rows in sheet order. Every column
of the index except `Sheet` is copied from the source row by header name; a column the source
sheet does not have (e.g. Acronym on Legislation) is left blank. Styling and row heights are
taken from the existing index rows, by parity, so the sheet keeps its appearance.

Usage:
    python3 rebuild_xref.py [workbook.xlsx ...]   # default: the canonical governance workbooks
                                                  #          in the working directory
Exit code 0 always; the caller decides what a change means. Prints what changed.
"""
import os
import sys
import openpyxl
from copy import copy

XREF = "Cross-Reference Index"
# Sheets that are not record sheets. A sheet without a Record_ID column is skipped anyway;
# this list is belt and braces for sheets that happen to have one.
NOT_RECORD_SHEETS = {XREF, "Data Dictionary", "Changelog", "Legend", "NCF Legend", "References"}


def source_records(wb):
    """[(record_id, {column: value}, sheet_name)] in workbook order, then sheet order."""
    out = []
    for sn in wb.sheetnames:
        if sn in NOT_RECORD_SHEETS:
            continue
        ws = wb[sn]
        first = next(ws.iter_rows(min_row=1, max_row=1), None)
        if not first:
            continue
        hdr = [c.value for c in first]
        if "Record_ID" not in hdr:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            rec = dict(zip(hdr, row))
            rid = rec.get("Record_ID")
            if rid and str(rid).strip():
                out.append((str(rid).strip(), rec, sn))
    return out


def rebuild(path, verbose=True):
    """Regenerate the index in `path`. Returns (n_rows, [change lines]); saves only if changed."""
    wb = openpyxl.load_workbook(path)
    if XREF not in wb.sheetnames:
        wb.close()
        return 0, []
    xr = wb[XREF]
    hdr = [c.value for c in next(xr.iter_rows(min_row=1, max_row=1))]
    ncol = len(hdr)

    before = {}
    for row in xr.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            before[str(row[0]).strip()] = dict(zip(hdr, row))
    old_max = xr.max_row

    # style templates, by row parity, captured before anything is overwritten
    tpl = {}
    for r in (2, 3):
        if r <= old_max:
            tpl[r % 2] = [
                {"font": copy(xr.cell(row=r, column=c).font),
                 "fill": copy(xr.cell(row=r, column=c).fill),
                 "border": copy(xr.cell(row=r, column=c).border),
                 "align": copy(xr.cell(row=r, column=c).alignment),
                 "nf": xr.cell(row=r, column=c).number_format}
                for c in range(1, ncol + 1)]
    height = xr.row_dimensions[2].height if 2 <= old_max else None

    records = source_records(wb)
    for i, (rid, rec, sn) in enumerate(records):
        r = i + 2
        for c, col in enumerate(hdr, start=1):
            cell = xr.cell(row=r, column=c)
            cell.value = sn if col == "Sheet" else rec.get(col)
            style = tpl.get(r % 2)
            if style and r > old_max:
                s = style[c - 1]
                cell.font, cell.fill, cell.border = copy(s["font"]), copy(s["fill"]), copy(s["border"])
                cell.alignment, cell.number_format = copy(s["align"]), s["nf"]
        if height and r > old_max:
            xr.row_dimensions[r].height = height

    last = len(records) + 1
    for r in range(last + 1, old_max + 1):          # clear any surplus trailing rows
        for c in range(1, ncol + 1):
            xr.cell(row=r, column=c).value = None

    after = {rid: {col: (sn if col == "Sheet" else rec.get(col)) for col in hdr}
             for rid, rec, sn in records}

    changes, per_col = [], {}
    added = [k for k in after if k not in before]
    removed = [k for k in before if k not in after]
    for rid in after:
        if rid in before:
            for col in hdr:
                if (before[rid].get(col) or None) != (after[rid].get(col) or None):
                    per_col[col] = per_col.get(col, 0) + 1
    if added:
        changes.append(f"added {len(added)} record(s) missing from the index: {', '.join(sorted(added))}")
    if removed:
        changes.append(f"removed {len(removed)} row(s) with no source record: {', '.join(sorted(removed))}")
    for col, n in sorted(per_col.items(), key=lambda x: -x[1]):
        changes.append(f"{col}: {n} value(s) refreshed from source")

    base = os.path.basename(path)
    if changes:
        wb.save(path)
        if verbose:
            print(f"  {base}: index rebuilt, {len(records)} rows")
            for line in changes:
                print(f"      - {line}")
    else:
        if verbose:
            print(f"  {base}: index already matches source ({len(records)} rows)")
    wb.close()
    return len(records), changes


def main():
    paths = sys.argv[1:] or [p for p in ("uk_climate_nature_governance.xlsx",
                                         "international_climate_nature_governance.xlsx")
                             if os.path.exists(p)]
    if not paths:
        print("No workbooks found. Run from Data/canonical_files/ or pass paths.")
        return 0
    any_change = False
    for p in paths:
        _, ch = rebuild(p)
        any_change = any_change or bool(ch)
    if any_change:
        print("\n  The index is a derived sheet, but it lives inside a canonical workbook: record "
              "this rebuild in the changelog entry for the version you are writing.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
